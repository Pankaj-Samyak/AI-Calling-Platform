import fitz
import docx
import io
import string
import random
import json
from livekit import api
import requests
from requests.auth import HTTPBasicAuth
from livekit.protocol.sip import CreateSIPOutboundTrunkRequest, SIPOutboundTrunkInfo
from dotenv import load_dotenv
load_dotenv()

#------------------------------User_Id-Generator----------------------------#
def generate_unique_id():
        characters = string.ascii_letters + string.digits
        return ''.join(random.choices(characters, k=8))

#-------------------------------to_snake_case-------------------------------#
def to_snake_case(col):
    return col.strip().replace(" ", "_").lower()
#-----------------------Campaign-Template-Validator-------------------------#
def validate_template(template: str, campaign_columns: list[str]) -> list[str]:
    errors = []
    try:
        # Check balanced braces
        open_braces = template.count('{')
        close_braces = template.count('}')
        if open_braces != close_braces:
              errors.append(f"Unmatched braces: {{ count = {open_braces} vs }} count = {close_braces}")

        # Extract placeholders
        formatter = string.Formatter()
        placeholders = {field_name for _, field_name, _, _ in formatter.parse(template) if field_name}

        # Find placeholders not in columns
        for ph in placeholders:
                if ph not in campaign_columns:
                        errors.append(f"Placeholder {{{ph}}} is not in campaign_columns")
        return errors
    except Exception as e:
        errors.append(f"error in validate_template function: {str(e)}")

#---------------------------------------------------------------------------#
# Extract text from in-memory .txt
def extract_text_from_txt(file_stream):
    return file_stream.read().decode('utf-8')

# Extract text from in-memory .pdf
def extract_text_from_pdf(file_stream):
    text = ""
    with fitz.open(stream=file_stream.read(), filetype="pdf") as doc:
        for page in doc:
            text += page.get_text()
    return text

# Extract text from in-memory .docx
def extract_text_from_docx(file_stream):
    doc = docx.Document(io.BytesIO(file_stream.read()))
    return " ".join([para.text for para in doc.paragraphs])

# Create Livekit Outbound SIP ID
async def get_lk_outbound_sip(name, address, numbers, user_name, password):
    livekit_api = api.LiveKitAPI()
    trunk = SIPOutboundTrunkInfo(
        name=name,
        address=address,
        numbers=[numbers],
        auth_username=user_name,
        auth_password=password,
    )
 
    request = CreateSIPOutboundTrunkRequest(trunk=trunk)
    created_trunk = await livekit_api.sip.create_sip_outbound_trunk(request)
    await livekit_api.aclose()
    return created_trunk.sip_trunk_id

# Trigger the call
# Generate a unique room name
async def trigger_outbound_call(outbound_trunk_id, system_prompt, campaign_id, batch_name, user_id):
    livekit_api = api.LiveKitAPI()
    # Generate a random room name for the outbound call
    # This can be adjusted to fit your naming conventions
    room_name = f"outbound-{''.join(str(random.randint(0, 9)) for _ in range(10))}"
    
    # Create the dispatch request
    request = api.CreateAgentDispatchRequest(
        agent_name="outbound-caller",
        room=room_name,
        metadata=json.dumps({
            "phone_number": "+919767288259",
            "outbound_trunk_id": outbound_trunk_id,
            "system_prompt": system_prompt,
            "campaign_id": campaign_id,
            "batch_name": batch_name,
            "user_id": user_id
        })
    )
    # Call the API
    await livekit_api.agent_dispatch.create_dispatch(request)
    print(f"Dispatch request created for room: {room_name}")

# Activate the call recording through twilio
def activate_recording(trunk_sid, account_sid, auth_token):
    try:
        url = f'https://trunking.twilio.com/v1/Trunks/{trunk_sid}/Recording'
        data = {
            'Mode': 'record-from-answer-dual'  # Options: 'do-not-record', 'record-from-ringing', 'record-from-answer'. For dual add -dual at end. 
        }
        response = requests.post(url, data=data, auth=HTTPBasicAuth(account_sid, auth_token))

        if response.status_code == 200:
            return {"status": "success", "message" : "Recording mode updated successfully."}
        else:
            return {"status": "fail", "message" : f"Failed to update recording mode. Status code: {response.status_code}"}
        
    except Exception as e:
         print(str(e), "Failure in activating the recording.")
         return {"status": "fail", "message" : "Failed to update recording mode, due to technical error."}